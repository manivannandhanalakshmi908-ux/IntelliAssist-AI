"""
IntelliAssist AI - Smart Document AI Assistant
RAG + hybrid search (FAISS embeddings + TF-IDF) + notes, summary, quiz,
sentiment & intent analysis, citations, evaluation.
Run:  streamlit run app.py
"""
import io
import os
import re
import json
import time
import hashlib
import requests
from concurrent.futures import ThreadPoolExecutor

import pandas as pd
import streamlit as st
from dotenv import load_dotenv
from pypdf import PdfReader
from docx import Document as DocxDocument
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import linear_kernel
from vaderSentiment.vaderSentiment import SentimentIntensityAnalyzer

from langchain_google_genai import ChatGoogleGenerativeAI
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.vectorstores import FAISS
from langchain_core.documents import Document

# ============================================================
# CONFIG
# ============================================================
st.set_page_config(page_title="IntelliAssist AI", page_icon="📚", layout="wide")
load_dotenv()

GOOGLE_API_KEY = os.getenv("GOOGLE_API_KEY")
if not GOOGLE_API_KEY:
    st.error("GOOGLE_API_KEY is missing from .env")
    st.stop()

DEFAULT_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.8-flash")
CHUNK_SIZE, CHUNK_OVERLAP = 900, 150
BATCH_CHARS = 14000      # text per map-reduce batch
MAX_BATCHES = 12         # cap LLM calls for very large documents

st.title("📚 IntelliAssist AI")
st.caption("Smart Document AI Assistant | Hybrid RAG • Notes • Insights • Citations")

# ============================================================
# SESSION STATE
# ============================================================
DEFAULTS = {
    "combined_hash": None, "file_names": [], "documents": [], "chunks": [],
    "db": None, "tfidf": None, "matrix": None, "insights": None,
    "chat_history": [], "summary": None, "notes": None, "quiz": None,
    "eval_df": None, "build_seconds": 0.0,
}
for key, value in DEFAULTS.items():
    st.session_state.setdefault(key, value)

@st.cache_data(ttl=600, show_spinner=False)
def list_models():
    """Ask Google which text-generation Flash models this API key can use."""
    try:
        r = requests.get("https://generativelanguage.googleapis.com/v1beta/models",
                         params={"key": GOOGLE_API_KEY, "pageSize": 200}, timeout=10)
        r.raise_for_status()
        skip = ("image", "live", "tts", "audio", "exp", "embedding", "robotics", "computer", "vision")
        names = []
        for m in r.json().get("models", []):
            n = m["name"].replace("models/", "")
            if ("generateContent" in m.get("supportedGenerationMethods", [])
                    and "flash" in n and not any(x in n for x in skip)):
                names.append(n)
        return sorted(set(names), reverse=True)
    except Exception:
        return []


# ============================================================
# SIDEBAR SETTINGS (needed before model load)
# ============================================================
with st.sidebar:
    st.header("📂 Upload Documents")
    uploaded_files = st.file_uploader(
        "PDF, DOCX or TXT (multiple allowed)",
        type=["pdf", "docx", "txt"],
        accept_multiple_files=True,
    )
    st.divider()
    st.subheader("⚙️ Settings")
    available_models = list_models()
    if available_models:
        default_idx = available_models.index(DEFAULT_MODEL) if DEFAULT_MODEL in available_models else 0
        model_name = st.selectbox("Gemini model", available_models, index=default_idx,
                                  help="Live list of models your API key can use.")
    else:
        model_name = st.text_input("Gemini model", DEFAULT_MODEL,
                                   help="Could not fetch the model list; type a model name.")
    fallback_models = st.text_input(
        "Extra fallback models (optional)", os.getenv("GEMINI_FALLBACKS", ""),
        help="Comma-separated. Other available models are also tried automatically on 503.")
    answer_style = st.selectbox("Answer detail", ["Concise", "Balanced", "Detailed"], index=2)
    top_k = st.slider("Chunks retrieved (k)", 3, 10, 5)
    use_hybrid = st.toggle("Hybrid search (Embeddings + TF-IDF)", value=True)

# ============================================================
# MODELS
# ============================================================
@st.cache_resource
def get_chat_model(name):
    return ChatGoogleGenerativeAI(model=name, google_api_key=GOOGLE_API_KEY,
                                  temperature=0.1, max_retries=1)

@st.cache_resource
def get_embeddings():
    return HuggingFaceEmbeddings(
        model_name="sentence-transformers/all-MiniLM-L6-v2",
        model_kwargs={"device": "cpu"},
        encode_kwargs={"normalize_embeddings": True, "batch_size": 64},
    )

@st.cache_resource
def get_vader():
    return SentimentIntensityAnalyzer()

try:
    chat_model = get_chat_model(model_name)
    embeddings = get_embeddings()
    vader = get_vader()
except Exception as e:
    st.error("Could not load AI models.")
    st.code(str(e))
    st.stop()


def to_text(content):
    """Gemini may return a str or a list of content parts."""
    if isinstance(content, str):
        return content
    if isinstance(content, list):
        return "".join(p if isinstance(p, str) else p.get("text", "") for p in content)
    return str(content)


TRANSIENT = ("503", "unavailable", "overloaded", "high demand", "429",
             "resource_exhausted", "500", "internal", "deadline", "timeout")
MISSING = ("404", "not found", "not supported")


def _is_transient(e):
    msg = str(e).lower()
    return any(t in msg for t in TRANSIENT)


def _is_missing(e):
    msg = str(e).lower()
    return any(t in msg for t in MISSING)


def model_chain():
    """Selected model, then manual fallbacks, then other models your key can use."""
    chain = [model_name.strip()]
    for m in fallback_models.split(","):
        m = m.strip()
        if m and m not in chain:
            chain.append(m)
    auto = sorted(available_models, key=lambda n: ("lite" not in n, n))
    for m in auto:
        if m not in chain:
            chain.append(m)
    return chain[:5]


def llm(prompt):
    """Retry on overload (503/429) with backoff, then fall back to the next model."""
    tried = []
    for name in model_chain():
        model = get_chat_model(name)
        for attempt in range(3):
            try:
                return to_text(model.invoke(prompt).content)
            except Exception as e:
                tried.append(f"{name}: {str(e)[:120]}")
                if _is_missing(e):
                    break                      # retired / bad model name -> next model
                if not _is_transient(e):
                    raise
                time.sleep(2 + attempt * 3)    # 2s, 5s, 8s
    raise RuntimeError("All models failed -> " + " | ".join(tried))


def llm_stream(prompt):
    """Streaming version with the same retry/fallback logic.
    Falls back only if the failure happens before any text was streamed."""
    tried = []
    for name in model_chain():
        model = get_chat_model(name)
        for attempt in range(3):
            started = False
            try:
                for chunk in model.stream(prompt):
                    text = to_text(chunk.content)
                    if text:
                        started = True
                        yield text
                return
            except Exception as e:
                tried.append(f"{name}: {str(e)[:120]}")
                if started:
                    raise
                if _is_missing(e):
                    break
                if not _is_transient(e):
                    raise
                time.sleep(2 + attempt * 3)
    raise RuntimeError("All models failed -> " + " | ".join(tried))

# ============================================================
# DOCUMENT READERS
# ============================================================
def paginate(name, text, size=3000):
    """DOCX/TXT have no pages: split into ~size-char sections ('Section N')."""
    paras = [p for p in text.split("\n") if p.strip()]
    sections, buf = [], ""
    for p in paras:
        if len(buf) + len(p) > size and buf:
            sections.append(buf)
            buf = ""
        buf += p + "\n"
    if buf.strip():
        sections.append(buf)
    return [Document(page_content=s.strip(), metadata={"source": name, "page": i})
            for i, s in enumerate(sections, start=1)]


def read_pdf(name, data):
    reader = PdfReader(io.BytesIO(data))
    docs = []
    for n, page in enumerate(reader.pages, start=1):
        text = page.extract_text()
        if text and text.strip():
            docs.append(Document(page_content=text.strip(),
                                 metadata={"source": name, "page": n}))
    return docs


def read_docx(name, data):
    doc = DocxDocument(io.BytesIO(data))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return paginate(name, "\n".join(parts))


def read_txt(name, data):
    return paginate(name, data.decode("utf-8", errors="ignore"))


def read_document(name, data):
    ext = name.lower().rsplit(".", 1)[-1]
    readers = {"pdf": read_pdf, "docx": read_docx, "txt": read_txt}
    if ext not in readers:
        raise ValueError("Only PDF, DOCX and TXT files are supported.")
    return readers[ext](name, data)


def split_documents(documents):
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE, chunk_overlap=CHUNK_OVERLAP,
        separators=["\n\n", "\n", ". ", "? ", "! ", " ", ""])
    chunks = [c for c in splitter.split_documents(documents) if len(c.page_content.strip()) > 40]
    for i, c in enumerate(chunks):
        c.metadata["chunk_id"] = i
    return chunks

# ============================================================
# INDEX (cached by file hash -> re-uploading the same files is instant)
# ============================================================
@st.cache_resource(show_spinner=False)
def build_index(key, _chunks):
    db = FAISS.from_documents(_chunks, embeddings)
    tfidf = TfidfVectorizer(stop_words="english", ngram_range=(1, 2),
                            sublinear_tf=True, max_features=50000)
    matrix = tfidf.fit_transform([c.page_content for c in _chunks])
    return db, tfidf, matrix


def compute_insights(documents, tfidf, matrix):
    full_text = " ".join(d.page_content for d in documents)
    words = len(full_text.split())
    terms = tfidf.get_feature_names_out()
    scores = matrix.sum(axis=0).A1
    top = scores.argsort()[::-1][:15]
    keywords = pd.DataFrame({"keyword": [terms[i] for i in top],
                             "score": [round(float(scores[i]), 3) for i in top]})
    rows = []
    for d in documents:
        comp = vader.polarity_scores(d.page_content[:5000])["compound"]
        rows.append({"file": d.metadata["source"], "page": d.metadata["page"], "sentiment": comp})
    sent_df = pd.DataFrame(rows)
    avg = float(sent_df["sentiment"].mean()) if len(sent_df) else 0.0
    return {"words": words, "reading_min": max(1, round(words / 220)),
            "keywords": keywords, "sentiment_df": sent_df, "avg_sentiment": avg}

# ============================================================
# RETRIEVAL (hybrid: FAISS semantic + TF-IDF lexical, fused with RRF)
# ============================================================
def retrieve(question, k):
    chunks = st.session_state.chunks
    pool = max(k * 2, 8)
    fused = {}

    sem = st.session_state.db.similarity_search(question, k=pool)
    for rank, d in enumerate(sem):
        cid = d.metadata["chunk_id"]
        fused[cid] = fused.get(cid, 0) + 1 / (60 + rank)

    if use_hybrid:
        qv = st.session_state.tfidf.transform([question])
        scores = linear_kernel(qv, st.session_state.matrix).ravel()
        order = [i for i in scores.argsort()[::-1][:pool] if scores[i] > 0]
        for rank, cid in enumerate(order):
            fused[cid] = fused.get(cid, 0) + 1 / (60 + rank)

    best = sorted(fused, key=fused.get, reverse=True)[:k]
    max_score = (2 if use_hybrid else 1) / 60
    return [(chunks[i], min(100, round(fused[i] / max_score * 100))) for i in best]


def make_context(results):
    return "\n".join(
        f"[SOURCE {n}] FILE: {d.metadata['source']} | PAGE: {d.metadata['page']}\n{d.page_content}\n"
        for n, (d, _) in enumerate(results, start=1))

# ============================================================
# SENTIMENT & INTENT ANALYSIS
# ============================================================
INTENT_RULES = [
    ("Summary", ["summar", "overview", "tl;dr", "gist", "in short"], "Give a structured summary with headings and bullets."),
    ("Comparison", ["difference", "compare", "versus", " vs ", "contrast", "distinguish"], "Use a markdown table comparing the items, then 1-2 lines of takeaway."),
    ("Procedure", ["how to", "steps", "procedure", "algorithm", "process of", "how do i"], "Give clearly numbered steps."),
    ("Definition", ["what is", "what are", "define", "meaning of", "definition"], "Start with a one-line definition, then key points and an example if present."),
    ("List", ["list", "enumerate", "types of", "name the", "kinds of", "examples of"], "Give a clear bulleted list with a short explanation for each item."),
    ("Explanation", ["explain", "why", "how does", "describe", "elaborate", "discuss"], "Explain step by step in simple language with bullets and bold key terms."),
    ("Factual", ["when", "who", "where", "how many", "how much", "which"], "Answer directly in 1-3 sentences."),
]


def detect_intent(question):
    q = " " + question.lower() + " "
    for name, keys, hint in INTENT_RULES:
        if any(k in q for k in keys):
            return name, hint
    return "General", "Answer clearly and completely using bullets where helpful."


def sentiment_label(score):
    if score >= 0.35:
        return "😊 Positive"
    if score <= -0.35:
        return "😟 Negative / frustrated"
    return "😐 Neutral"

# ============================================================
# PROMPTS
# ============================================================
STYLE_RULES = {
    "Concise": "Keep the answer short (under 120 words).",
    "Balanced": "Give a clear answer of moderate length.",
    "Detailed": "Give a thorough, well-organised answer: short intro, detailed bullet points with "
                "definitions/examples from the document, and a one-line takeaway.",
}


def history_text(n=6):
    msgs = st.session_state.chat_history[-n:]
    return "\n".join(f"{m['role'].upper()}: {m['content'][:500]}" for m in msgs)


def rewrite_question(question):
    """Turn follow-ups like 'explain it more' into standalone search queries."""
    if not st.session_state.chat_history or len(question.split()) > 12:
        return question
    try:
        out = llm(f"""Rewrite the follow-up question as a standalone question using the chat history.
Return ONLY the rewritten question.

HISTORY:
{history_text(4)}

FOLLOW-UP: {question}""").strip()
        return out or question
    except Exception:
        return question


def build_answer_prompt(question, standalone, results, intent_hint):
    return f"""You are IntelliAssist AI, an expert study assistant. Answer ONLY from the DOCUMENT CONTEXT.

DOCUMENT CONTEXT:
{make_context(results)}

RECENT CHAT:
{history_text()}

QUESTION: {question}
(Standalone form: {standalone})

RULES:
1. Answer the exact question; no outside knowledge; never invent facts or pages.
2. FORMAT: {intent_hint}
3. DETAIL: {STYLE_RULES[answer_style]}
4. Cite inline like [Page 3] after the facts they support. Only cite pages in the context.
5. If the context lacks the answer, say exactly: "I could not find enough information to answer this question in the uploaded document."
6. End with a line: 📌 Sources: <file> - Pages X, Y
"""


def answer_blocking(question):
    results = retrieve(question, top_k)
    intent, hint = detect_intent(question)
    return llm(build_answer_prompt(question, question, results, hint)), results

# ============================================================
# MAP-REDUCE HELPERS (parallel LLM calls = fast on big documents)
# ============================================================
def make_batches(documents):
    batches, buf = [], ""
    for d in documents:
        piece = f"\n[{d.metadata['source']} - Page {d.metadata['page']}]\n{d.page_content}\n"
        if len(buf) + len(piece) > BATCH_CHARS and buf:
            batches.append(buf)
            buf = ""
        buf += piece
    if buf.strip():
        batches.append(buf)
    if len(batches) > MAX_BATCHES:  # sample evenly to bound cost
        step = len(batches) / MAX_BATCHES
        batches = [batches[int(i * step)] for i in range(MAX_BATCHES)]
    return batches


def parallel_llm(prompts, workers=4):
    with ThreadPoolExecutor(max_workers=workers) as ex:
        return list(ex.map(llm, prompts))


def generate_notes():
    batches = make_batches(st.session_state.documents)
    prompts = [f"""You are an expert teacher creating DETAILED STUDY NOTES from the material below.
Use ONLY this material. Include page references like (p. 4).

MATERIAL:
{b}

Produce markdown with:
## <Topic heading> for each major topic
- Clear explanation bullets (not just keywords)
- **Key terms** in bold with one-line definitions
- Formulas, steps, examples, or diagrams described in the text
- ⚠️ Important points / common mistakes if present
Be thorough but organised.""" for b in batches]
    parts = parallel_llm(prompts)
    notes = "\n\n".join(parts)
    final = llm(f"""From these study notes, write (markdown):
## 🎯 Key Takeaways (6-8 bullets)
## 📖 Glossary (10-15 important terms with definitions)
## ❓ Likely Exam Questions (8 questions, mix of short and long answer)
## ⚡ Quick Revision Sheet (10 one-line points)

NOTES:
{notes[:30000]}""")
    return f"# Study Notes\n\n{final}\n\n---\n\n# Detailed Notes\n\n{notes}"


def generate_summary(mode):
    spec = {"Brief": "6-8 bullet points total.",
            "Standard": "Headings per major topic with 3-5 bullets each and a short overview.",
            "Detailed": "Overview, then headings per topic with detailed bullets, key concepts, "
                        "important facts/figures, and a conclusion."}[mode]
    docs = st.session_state.documents
    total = sum(len(d.page_content) for d in docs)
    if total <= BATCH_CHARS * 2:
        text = "\n".join(f"[{d.metadata['source']} - Page {d.metadata['page']}]\n{d.page_content}" for d in docs)
        return llm(f"""Summarize the document using ONLY its content. Student-friendly. Mention pages.
FORMAT: {spec}

DOCUMENT:
{text}""")
    parts = parallel_llm([f"Summarize this part of a document faithfully with page references, in 10-15 bullets:\n{b}"
                          for b in make_batches(docs)])
    return llm(f"""Combine these partial summaries into one coherent summary. Use ONLY them. Mention pages.
FORMAT: {spec}

PARTS:
{chr(10).join(parts)}""")


def generate_quiz(n):
    sample = "\n".join(c.page_content for c in st.session_state.chunks[:60])[:20000]
    raw = llm(f"""Create {n} multiple-choice questions from the text. Return ONLY JSON:
[{{"question":"...","options":["A","B","C","D"],"answer_index":0,"explanation":"..."}}]

TEXT:
{sample}""")
    raw = re.sub(r"^```(?:json)?|```$", "", raw.strip(), flags=re.M).strip()
    return json.loads(raw)

# ============================================================
# EXPORT HELPERS
# ============================================================
def markdown_to_docx(md, title):
    doc = DocxDocument()
    doc.add_heading(title, 0)
    for line in md.splitlines():
        s = line.rstrip()
        if not s.strip() or s.strip() == "---":
            continue
        clean = s.replace("**", "")
        if s.startswith("### "):
            doc.add_heading(clean[4:], 3)
        elif s.startswith("## "):
            doc.add_heading(clean[3:], 2)
        elif s.startswith("# "):
            doc.add_heading(clean[2:], 1)
        elif re.match(r"^\s*[-*•]\s+", s):
            doc.add_paragraph(re.sub(r"^\s*[-*•]\s+", "", clean), style="List Bullet")
        elif re.match(r"^\s*\d+[.)]\s+", s):
            doc.add_paragraph(re.sub(r"^\s*\d+[.)]\s+", "", clean), style="List Number")
        else:
            doc.add_paragraph(clean)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def chat_to_markdown():
    lines = ["# IntelliAssist AI - Conversation\n"]
    for m in st.session_state.chat_history:
        who = "🧑 You" if m["role"] == "user" else "🤖 IntelliAssist"
        lines.append(f"**{who}:**\n\n{m['content']}\n")
    return "\n".join(lines)

# ============================================================
# PROCESS UPLOADS
# ============================================================
if uploaded_files:
    payload = [(f.name, f.getvalue()) for f in uploaded_files]
    combined = hashlib.md5("".join(hashlib.md5(d).hexdigest() for _, d in payload).encode()).hexdigest()

    if combined != st.session_state.combined_hash:
        start = time.perf_counter()
        with st.status("📖 Processing documents...", expanded=True) as status:
            documents = []
            for name, data in payload:
                st.write(f"Reading {name}...")
                try:
                    documents += read_document(name, data)
                except Exception as e:
                    status.update(label=f"❌ Failed to read {name}", state="error")
                    st.error(str(e))
                    st.stop()
            if not documents:
                status.update(label="❌ No text found", state="error")
                st.error("No readable text found (scanned PDFs need OCR).")
                st.stop()

            st.write("Splitting into chunks...")
            chunks = split_documents(documents)
            st.write(f"Building hybrid index for {len(chunks)} chunks...")
            try:
                db, tfidf, matrix = build_index(combined, chunks)
            except Exception as e:
                status.update(label="❌ Index creation failed", state="error")
                st.error(str(e))
                st.stop()

            st.write("Analysing sentiment & keywords...")
            insights = compute_insights(documents, tfidf, matrix)
            status.update(label="✅ Documents ready!", state="complete")

        st.session_state.update(
            combined_hash=combined, file_names=[n for n, _ in payload],
            documents=documents, chunks=chunks, db=db, tfidf=tfidf, matrix=matrix,
            insights=insights, chat_history=[], summary=None, notes=None,
            quiz=None, eval_df=None, build_seconds=time.perf_counter() - start)

# ============================================================
# SIDEBAR STATUS
# ============================================================
with st.sidebar:
    st.divider()
    st.subheader("📊 Document Status")
    if st.session_state.file_names:
        st.write(f"**Files:** {', '.join(st.session_state.file_names)}")
        st.write(f"**Pages/Sections:** {len(st.session_state.documents)}")
        st.write(f"**Chunks:** {len(st.session_state.chunks)}")
        st.write(f"**Indexed in:** {st.session_state.build_seconds:.1f}s")
        if st.session_state.chat_history:
            st.download_button("💾 Export chat (.md)", chat_to_markdown(),
                               "intelliassist_chat.md", use_container_width=True)
            if st.button("🗑️ Clear chat", use_container_width=True):
                st.session_state.chat_history = []
                st.rerun()
    else:
        st.write("No document uploaded.")


def render_sources(sources):
    for s in sources:
        st.markdown(f"**{s['file']} — Page {s['page']}** · relevance {s['score']}%")
        st.caption(s["preview"])

# ============================================================
# MAIN APP
# ============================================================
if st.session_state.db:
    tabs = st.tabs(["💬 Chat", "📝 Study Notes", "📋 Summary", "📈 Insights",
                    "🧩 Quiz", "🧪 Evaluation", "ℹ️ Project"])
    tab_chat, tab_notes, tab_summary, tab_insights, tab_quiz, tab_eval, tab_info = tabs

    # ---------------- CHAT ----------------
    with tab_chat:
        for m in st.session_state.chat_history:
            with st.chat_message(m["role"]):
                st.markdown(m["content"])
                if m["role"] == "user" and m.get("analysis"):
                    st.caption(m["analysis"])
                if m["role"] == "assistant":
                    if m.get("meta"):
                        st.caption(m["meta"])
                    if m.get("sources"):
                        with st.expander("📌 Retrieved Sources"):
                            render_sources(m["sources"])

        question = st.chat_input("Ask something about your document...")
        if question:
            intent, hint = detect_intent(question)
            mood = sentiment_label(vader.polarity_scores(question)["compound"])
            analysis = f"Intent: **{intent}** · Tone: {mood}"
            if mood.startswith("😟"):
                hint += " The user sounds frustrated: be extra clear, patient and reassuring."

            with st.chat_message("user"):
                st.markdown(question)
                st.caption(analysis)

            with st.chat_message("assistant"):
                t0 = time.perf_counter()
                try:
                    with st.spinner("🔎 Searching..."):
                        standalone = rewrite_question(question)
                        results = retrieve(standalone, top_k)
                    prompt = build_answer_prompt(question, standalone, results, hint)
                    answer = st.write_stream(llm_stream(prompt))
                except Exception as e:
                    st.error("Could not generate the answer.")
                    st.code(str(e))
                    st.stop()
                elapsed = time.perf_counter() - t0

                sources, seen = [], set()
                for d, score in results:
                    preview = d.page_content.replace("\n", " ").strip()
                    preview = preview[:300] + ("..." if len(preview) > 300 else "")
                    key = (d.metadata["source"], d.metadata["page"], preview[:60])
                    if key not in seen:
                        seen.add(key)
                        sources.append({"file": d.metadata["source"], "page": d.metadata["page"],
                                        "score": score, "preview": preview})
                meta = f"⏱️ {elapsed:.1f}s · {len(results)} chunks · {'hybrid' if use_hybrid else 'semantic'} search"
                st.caption(meta)
                with st.expander("📌 Retrieved Sources"):
                    render_sources(sources)

            st.session_state.chat_history.append(
                {"role": "user", "content": question, "analysis": analysis})
            st.session_state.chat_history.append(
                {"role": "assistant", "content": answer, "sources": sources, "meta": meta})

    # ---------------- NOTES ----------------
    with tab_notes:
        st.subheader("📝 Detailed Study Notes")
        st.caption("Topic-wise notes, key terms, glossary, exam questions and a revision sheet. "
                   "Large documents are processed in parallel for speed.")
        if st.button("✨ Generate Study Notes", use_container_width=True):
            with st.spinner("Writing detailed notes..."):
                try:
                    t0 = time.perf_counter()
                    st.session_state.notes = generate_notes()
                    st.success(f"Done in {time.perf_counter() - t0:.1f}s")
                except Exception as e:
                    st.error("Notes generation failed.")
                    st.code(str(e))
        if st.session_state.notes:
            c1, c2 = st.columns(2)
            c1.download_button("⬇️ Download .md", st.session_state.notes,
                               "study_notes.md", use_container_width=True)
            c2.download_button("⬇️ Download .docx",
                               markdown_to_docx(st.session_state.notes, "IntelliAssist Study Notes"),
                               "study_notes.docx", use_container_width=True)
            st.markdown(st.session_state.notes)
        else:
            st.info("Click the button to generate notes.")

    # ---------------- SUMMARY ----------------
    with tab_summary:
        st.subheader("📋 Document Summary")
        mode = st.radio("Length", ["Brief", "Standard", "Detailed"], index=1, horizontal=True)
        if st.button("✨ Generate Summary", use_container_width=True):
            with st.spinner("Creating summary..."):
                try:
                    st.session_state.summary = generate_summary(mode)
                except Exception as e:
                    st.error("Summary generation failed.")
                    st.code(str(e))
        if st.session_state.summary:
            st.download_button("⬇️ Download summary", st.session_state.summary, "summary.md")
            st.markdown(st.session_state.summary)
        else:
            st.info("Choose a length and click generate.")

    # ---------------- INSIGHTS ----------------
    with tab_insights:
        ins = st.session_state.insights
        st.subheader("📈 Document Insights")
        c1, c2, c3, c4 = st.columns(4)
        c1.metric("Words", f"{ins['words']:,}")
        c2.metric("Reading time", f"{ins['reading_min']} min")
        c3.metric("Pages/Sections", len(st.session_state.documents))
        c4.metric("Overall sentiment", sentiment_label(ins["avg_sentiment"]),
                  f"{ins['avg_sentiment']:+.2f}")
        left, right = st.columns(2)
        with left:
            st.markdown("**🔑 Top keywords (TF-IDF)**")
            st.bar_chart(ins["keywords"].set_index("keyword"))
        with right:
            st.markdown("**🎭 Sentiment by page (VADER)**")
            sdf = ins["sentiment_df"].copy()
            sdf["label"] = sdf["file"].str[:12] + " p" + sdf["page"].astype(str)
            st.line_chart(sdf.set_index("label")["sentiment"])
        st.caption("Sentiment ranges from -1 (negative) to +1 (positive). "
                   "Technical documents are usually near neutral.")

    # ---------------- QUIZ ----------------
    with tab_quiz:
        st.subheader("🧩 Practice Quiz")
        n_q = st.slider("Number of questions", 3, 10, 5)
        if st.button("🎲 Generate Quiz", use_container_width=True):
            with st.spinner("Creating quiz..."):
                try:
                    st.session_state.quiz = generate_quiz(n_q)
                except Exception as e:
                    st.error("Quiz generation failed. Try again.")
                    st.code(str(e))
        for i, q in enumerate(st.session_state.quiz or []):
            st.markdown(f"**Q{i + 1}. {q['question']}**")
            choice = st.radio("Choose", q["options"], index=None, key=f"quiz_{i}",
                              label_visibility="collapsed")
            if choice is not None:
                correct = q["options"][q["answer_index"]]
                if choice == correct:
                    st.success(f"✅ Correct! {q.get('explanation', '')}")
                else:
                    st.error(f"❌ Answer: {correct}. {q.get('explanation', '')}")

    # ---------------- EVALUATION ----------------
    with tab_eval:
        st.subheader("🧪 Evaluation (for your report)")
        st.caption("One test per line:  `question || keyword1, keyword2`  — the app measures "
                   "latency, keyword recall in the answer and in retrieved chunks.")
        sample = st.text_area("Test cases", height=160,
                              placeholder="What is overfitting? || training data, generalize")
        if st.button("▶️ Run evaluation", use_container_width=True) and sample.strip():
            rows = []
            bar = st.progress(0.0)
            lines = [l for l in sample.splitlines() if l.strip()]
            for i, line in enumerate(lines):
                q, _, kws = line.partition("||")
                keywords = [k.strip().lower() for k in kws.split(",") if k.strip()]
                t0 = time.perf_counter()
                try:
                    ans, res = answer_blocking(q.strip())
                except Exception as e:
                    ans, res = f"ERROR: {e}", []
                lat = time.perf_counter() - t0
                ctx = " ".join(d.page_content for d, _ in res).lower()
                hit = lambda text: (sum(k in text.lower() for k in keywords) / len(keywords)) if keywords else None
                rows.append({"question": q.strip(), "latency_s": round(lat, 2),
                             "answer_keyword_recall": hit(ans), "retrieval_keyword_recall": hit(ctx),
                             "answered": "could not find enough information" not in ans.lower(),
                             "answer": ans[:300]})
                bar.progress((i + 1) / len(lines))
            st.session_state.eval_df = pd.DataFrame(rows)
        if st.session_state.eval_df is not None:
            df = st.session_state.eval_df
            c1, c2, c3 = st.columns(3)
            c1.metric("Avg latency", f"{df['latency_s'].mean():.2f}s")
            c2.metric("Avg answer recall", f"{df['answer_keyword_recall'].dropna().mean():.0%}"
                      if df["answer_keyword_recall"].notna().any() else "n/a")
            c3.metric("Answered", f"{df['answered'].mean():.0%}")
            st.dataframe(df, use_container_width=True)
            st.download_button("⬇️ Download CSV", df.to_csv(index=False), "evaluation.csv")

    # ---------------- PROJECT ----------------
    with tab_info:
        st.subheader("🚀 IntelliAssist AI")
        st.markdown("""
**Problem:** Students and businesses struggle to search across PDFs, notes and documents.
**Solution:** A RAG chatbot over your own documents, with notes, summaries, insights and citations.

**Pipeline:** Upload → Extract → Chunk → MiniLM embeddings (BERT-family) + TF-IDF → FAISS →
Hybrid retrieval (Reciprocal Rank Fusion) → Gemini (streaming) → Answer + Sources

**Features:** PDF/DOCX/TXT (multi-file) · hybrid semantic search · history-aware chat ·
intent & sentiment analysis · detailed study notes (.md/.docx) · 3 summary lengths ·
keyword & sentiment insights · quiz · evaluation harness · chat export

**Stack:** Python, Streamlit, LangChain, Hugging Face, FAISS, scikit-learn, VADER, Google Gemini
""")

else:
    st.info("👈 Upload one or more PDF, DOCX or TXT files to start.")
    st.markdown("""
## 📚 IntelliAssist AI
Upload study material, notes or reports and:

💬 **Chat** with page-level citations · 📝 **Generate detailed study notes** ·
📋 **Summarise** · 📈 **See keywords & sentiment** · 🧩 **Take a quiz** · 🧪 **Evaluate accuracy**
""")